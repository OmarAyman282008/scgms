import io
import os

from django.conf import settings
from django.utils import timezone


def to_xlsx(rep, rtl):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = rep.slug[:30]
    ws.sheet_view.rightToLeft = rtl
    ws.append([rep.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([rep.period])
    ws.append([])
    for k, v in rep.summary:
        ws.append([k, v])
    if rep.summary:
        ws.append([])
    ws.append(rep.columns)
    hdr = ws.max_row
    for cell in ws[hdr]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3B6E")
    for r in rep.rows:
        ws.append(list(r))
    for col in ws.columns:
        width = min(60, max(len(str(c.value or "")) for c in col[hdr - 1:]) + 2)
        ws.column_dimensions[col[0].column_letter].width = max(10, width)
        for c in col[hdr:]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    if rep.note:
        ws.append([])
        ws.append([rep.note])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _shape(text, rtl):
    text = str(text)
    if not any("؀" <= ch <= "ۿ" for ch in text):
        return text
    import arabic_reshaper
    from bidi.algorithm import get_display
    return get_display(arabic_reshaper.reshape(text))


def to_pdf(rep, rtl):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fdir = os.path.join(settings.BASE_DIR, "static", "fonts")
    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", os.path.join(fdir, "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", os.path.join(fdir, "DejaVuSans-Bold.ttf")))
    align = 2 if rtl else 0
    st = ParagraphStyle("b", fontName="DejaVu", fontSize=7.5, leading=9.5, alignment=align)
    sh = ParagraphStyle("h", fontName="DejaVu-Bold", fontSize=7.5, leading=9.5, alignment=align, textColor=colors.white)
    st_title = ParagraphStyle("t", fontName="DejaVu-Bold", fontSize=14, leading=18, alignment=align)
    st_sub = ParagraphStyle("s", fontName="DejaVu", fontSize=9, leading=12, alignment=align, textColor=colors.HexColor("#555555"))
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.2 * cm,
                            bottomMargin=1.2 * cm, title=rep.title)
    els = []
    logo = os.path.join(settings.BASE_DIR, "static", "img", "logo.jpg")
    head = [[Paragraph(_shape(rep.title, rtl), st_title)],
            [Paragraph(_shape(f"{settings.SCHOOL_NAME_AR if rtl else settings.SCHOOL_NAME_EN} · {rep.period}", rtl), st_sub)]]
    ht = Table([[Image(logo, 1.6 * cm, 1.6 * cm), Table(head)]] if not rtl else [[Table(head), Image(logo, 1.6 * cm, 1.6 * cm)]],
               colWidths=[2 * cm, page[0] - 4.4 * cm] if not rtl else [page[0] - 4.4 * cm, 2 * cm])
    els += [ht, Spacer(1, 0.3 * cm)]
    if rep.summary:
        srow = [Paragraph(_shape(f"{k}: {v}", rtl), st) for k, v in rep.summary]
        if rtl:
            srow = srow[::-1]
        stbl = Table([srow])
        stbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF2F8")),
                                  ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#C9D3E3"))]))
        els += [stbl, Spacer(1, 0.3 * cm)]
    cols = rep.columns[::-1] if rtl else rep.columns
    data = [[Paragraph(_shape(c, rtl), sh) for c in cols]]
    for r in rep.rows:
        r = list(r)[::-1] if rtl else list(r)
        data.append([Paragraph(_shape(v, rtl), st) for v in r])
    if len(data) == 1:
        data.append([Paragraph("—", st)] + [""] * (len(cols) - 1))
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3B6E")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D3E3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FB")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    els.append(t)
    if rep.note:
        els += [Spacer(1, 0.3 * cm), Paragraph(_shape(rep.note, rtl), st_sub)]
    els += [Spacer(1, 0.3 * cm), Paragraph(_shape(f"{timezone.localtime():%Y-%m-%d %H:%M} · SCGMS", rtl), st_sub)]
    doc.build(els)
    buf.seek(0)
    return buf
