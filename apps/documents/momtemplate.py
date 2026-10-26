"""Fill the approved MOM template for a meeting.

Any uploaded DOCX template may use these placeholders (each typed in one go, so
Word keeps it in one piece). Missing placeholders are simply skipped, so a plain
template without placeholders downloads unchanged.

  Inline:  {{committee_name}} {{meeting_number}} {{year}} {{date}} {{day}} {{time}}
           {{venue}} {{meeting_type}} {{chair}} {{secretary}} {{absentees}} {{issue_date}}
  Blocks (alone in their paragraph):
           {{attendees_table}} {{agenda}} {{previous_actions_table}} {{discussion}}
           {{decisions_table}}
"""
import copy
import io
import os
import re

from django.conf import settings
from docx import Document as Docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import DocumentTemplate

DEFAULT_TEMPLATE = os.path.join(os.path.dirname(__file__), "default_templates", "mom_template.docx")
PH = re.compile(r"\{\{\s*([a-z_]+)\s*\}\}")
DAYS_AR = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
ROLE_AR = {"chair": "رئيس اللجنة", "secretary": "مقرر اللجنة", "member": "عضو"}
STATUS_AR = {"registered": "مسجل", "assigned": "تم الإسناد", "in_progress": "قيد التنفيذ",
             "implemented": "منفذ ومتحقق منه", "noted": "للعلم", "cancelled": "ملغى", "superseded": "مستبدل"}
KIND_AR = {"regular": "اجتماع دوري", "exceptional": "اجتماع استثنائي", "urgent": "اجتماع عاجل"}
DOTS = "." * 110


def template_path():
    t = DocumentTemplate.active("mom")
    if t and t.file and os.path.exists(t.file.path):
        return t.file.path
    return DEFAULT_TEMPLATE


def _rtl_run(run, bold=False, size=11):
    run.bold = bold
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    rpr.append(OxmlElement("w:rtl"))
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(size * 2))
    rpr.append(szcs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))


def _bidi(p):
    p._p.get_or_add_pPr().insert(0, OxmlElement("w:bidi"))


def _table(doc, anchor_p, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    tblpr = t._tbl.tblPr
    tblpr.append(OxmlElement("w:bidiVisual"))
    borders = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{e}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "808080")):
            b.set(qn(f"w:{k}"), v)
        borders.append(b)
    tblpr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 567)))

    def fill(cell, text, bold=False, shade=None):
        cell.text = ""
        p = cell.paragraphs[0]
        _bidi(p)
        _rtl_run(p.add_run(str(text)), bold=bold, size=10)
        if shade:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), shade)
            cell._tc.get_or_add_tcPr().append(shd)

    for i, h in enumerate(headers):
        fill(t.rows[0].cells[i], h, True, "E4EAF3")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            fill(cells[i], v)
    for row in t.rows:
        row.height = Cm(0.7)
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    anchor_p._p.addnext(t._tbl)
    return t


def _lines(doc, anchor_p, lines):
    prev = anchor_p._p
    for line in lines:
        newp = copy.deepcopy(anchor_p._p)
        for r in newp.findall(qn("w:r")):
            newp.remove(r)
        prev.addnext(newp)
        prev = newp
        from docx.text.paragraph import Paragraph
        p = Paragraph(newp, anchor_p._parent)
        _rtl_run(p.add_run(line), size=11)


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in doc.sections:
        for part in (s.header, s.first_page_header, s.footer):
            if part is not None:
                yield from part.paragraphs


def _replace_inline(p, values):
    sub = lambda t: PH.sub(lambda m: str(values.get(m.group(1), m.group(0))), t)  # noqa: E731
    for r in p.runs:  # placeholders typed in one go sit inside one run: keep its formatting
        if "{{" in r.text:
            r.text = sub(r.text)
    runs = p.runs
    text = "".join(r.text for r in runs)
    if runs and PH.search(text):  # a placeholder split across runs: rewrite the paragraph text
        new = sub(text)
        if new != text:
            runs[0].text = new
            for r in runs[1:]:
                r.text = ""


def build_context(meeting):
    from apps.governance.models import Decision
    c = meeting.committee
    members = list(c.active_memberships())
    att = {a.user_id: a.status for a in meeting.attendance.all()}
    order = {"chair": 0, "secretary": 1, "member": 2}
    members.sort(key=lambda m: (order.get(m.role, 3), m.order))
    attendees = [m for m in members if att.get(m.user_id, "present") == "present"]
    absent = [m for m in members if att.get(m.user_id) in ("absent", "excused")]
    prev = (Decision.objects.filter(committee=c, status__in=Decision.OPEN_STATUSES, meeting__date__lt=meeting.date)
            .select_related("responsible").order_by("meeting__date", "id"))
    values = {
        "committee_name": c.name_ar,
        "meeting_number": meeting.number,
        "year": meeting.date.year,
        "date": meeting.date.strftime("%d/%m/%Y"),
        "day": DAYS_AR[meeting.date.weekday()],
        "time": meeting.start_time.strftime("%H:%M") if meeting.start_time else "..........",
        "venue": meeting.venue or "....................",
        "meeting_type": KIND_AR.get(meeting.kind, meeting.kind),
        "chair": c.chair.name_ar if c.chair else "....................",
        "secretary": c.secretary.name_ar if c.secretary else "....................",
        "absentees": "، ".join(f"{m.user.name_ar or m.user.username} ({'بعذر' if att.get(m.user_id) == 'excused' else 'غياب'})"
                               for m in absent) or "لا يوجد",
        "issue_date": meeting.date.strftime("%d/%m/%Y"),
    }
    blocks = {
        "attendees": [(str(i), m.user.name_ar or m.user.username, ROLE_AR.get(m.role, m.role), "")
                      for i, m in enumerate(attendees, 1)] + [(str(len(attendees) + 1), "", "ضيف / مدعو", "")],
        "agenda": [f"{i}- {a.title}" for i, a in enumerate(meeting.agenda_items.all(), 1)] or
                  [f"{i}- {DOTS}" for i in range(1, 4)],
        "previous": [(d.code, d.title, d.responsible.name_ar if d.responsible else "-",
                      d.target_date.strftime("%d/%m/%Y") if d.target_date else "-", STATUS_AR.get(d.status, d.status))
                     for d in prev],
        "discussion": (meeting.mom.discussion_summary.splitlines() if meeting.mom.discussion_summary else
                       [f"{i}- {DOTS}" for i in range(1, 4)]),
    }
    return values, blocks


def render(meeting):
    doc = Docx(template_path())
    values, blocks = build_context(meeting)
    for p in list(_all_paragraphs(doc)):
        text = "".join(r.text for r in p.runs).strip()
        m = PH.fullmatch(text)
        key = m.group(1) if m else None
        if key == "attendees_table":
            _table(doc, p, ["م", "الاسم", "الصفة", "التوقيع"], blocks["attendees"], [1.0, 5.6, 4.6, 4.2])
        elif key == "previous_actions_table":
            if blocks["previous"]:
                _table(doc, p, ["رقم القرار", "القرار", "المسؤول", "الموعد", "الحالة"], blocks["previous"],
                       [3.0, 6.4, 2.8, 2.0, 2.0])
            else:
                _lines(doc, p, ["لا توجد قرارات سابقة مفتوحة."])
        elif key == "decisions_table":
            _table(doc, p, ["م", "القرار / التوصية", "قابل للتنفيذ / للعلم", "المسؤول", "الموعد",
                            "الشواهد المطلوبة"], [(str(i), "", "", "", "", "") for i in range(1, 7)],
                   [0.9, 6.0, 2.2, 2.6, 1.8, 2.7])
        elif key == "agenda":
            _lines(doc, p, blocks["agenda"])
        elif key == "discussion":
            _lines(doc, p, blocks["discussion"])
        else:
            _replace_inline(p, values)
            continue
        p._p.getparent().remove(p._p)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
